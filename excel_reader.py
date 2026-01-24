from openpyxl import load_workbook


def read_test_data(file_path):
    try:
        workbook = load_workbook(file_path)
        sheet = workbook.active
        test_data = {}
        
        # Column A: Field Name, Column B: Value
        for row in range(2, sheet.max_row + 1):
            field_name = sheet.cell(row=row, column=1).value
            field_value = sheet.cell(row=row, column=2).value
            
            if field_name and field_value is not None:
                test_data[field_name] = field_value
        
        workbook.close()
        
        print(f"Successfully loaded test data from: {file_path}")
        print(f"Total fields loaded: {len(test_data)}")
        
        return test_data
        
    except FileNotFoundError:
        print(f"Error: Excel file not found at {file_path}")
        raise
    except Exception as e:
        print(f"Error reading Excel file: {e}")
        raise

if __name__ == "__main__":
    # Test the excel reader
    print("Testing Excel Reader...")
    print("\nMake sure you have created 'test_data.xlsx' file with proper structure.")
    print("\nExpected Excel structure:")
    print("Column A: Field Name | Column B: Value")
    print("Example:")
    print("url | https://www.aviasales.kz")
    print("from_city | Астана")
    print("to_city | Уральск")
