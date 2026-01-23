"""
Excel File Generator
This script creates the test_data.xlsx file with sample data.
Run this script to generate the Excel file needed for the tests.

Usage:
    python create_excel_file.py
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import os


def create_test_data_excel():
    """
    Creates test_data.xlsx with test data and BrowserStack configuration.
    """
    print("="*60)
    print("Excel File Generator")
    print("="*60)
    
    # Create a new workbook
    wb = Workbook()
    
    # =========================================
    # SHEET 1: TestData
    # =========================================
    ws_test = wb.active
    ws_test.title = "TestData"
    
    print("\n[1/2] Creating TestData sheet...")
    
    # Style for headers
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    
    # Add headers
    ws_test['A1'] = 'Field Name'
    ws_test['B1'] = 'Value'
    
    # Apply header styling
    for cell in ['A1', 'B1']:
        ws_test[cell].fill = header_fill
        ws_test[cell].font = header_font
        ws_test[cell].alignment = Alignment(horizontal='center', vertical='center')
    
    # Test data
    test_data = [
        ['url', 'https://www.aviasales.kz'],
        ['from_city', 'Астана'],
        ['to_city', 'Уральск'],
        ['email', '23232323@astanait.edu.kz'],
        ['phone', '7787296919'],
        ['name', 'Meruyert'],
        ['lastname', 'Boranbay'],
        ['birth_day', 2],
        ['birth_month', 7],
        ['birth_year', 2005],
        ['passport_number', '123456798'],
        ['passport_exp_day', 12],
        ['passport_exp_month', 12],
        ['passport_exp_year', 2030],
        ['nationality', 'Казахстан']
    ]
    
    # Add test data
    for idx, row in enumerate(test_data, start=2):
        ws_test[f'A{idx}'] = row[0]
        ws_test[f'B{idx}'] = row[1]
    
    # Adjust column widths
    ws_test.column_dimensions['A'].width = 25
    ws_test.column_dimensions['B'].width = 35
    
    print(f"✓ Added {len(test_data)} test data fields")
    
    # =========================================
    # SHEET 2: BrowserStack
    # =========================================
    ws_bs = wb.create_sheet(title="BrowserStack")
    
    print("\n[2/2] Creating BrowserStack configuration sheet...")
    
    # Add headers
    ws_bs['A1'] = 'Field Name'
    ws_bs['B1'] = 'Value'
    
    # Apply header styling
    for cell in ['A1', 'B1']:
        ws_bs[cell].fill = header_fill
        ws_bs[cell].font = header_font
        ws_bs[cell].alignment = Alignment(horizontal='center', vertical='center')
    
    # BrowserStack configuration
    bs_data = [
        ['username', 'YOUR_BROWSERSTACK_USERNAME'],
        ['access_key', 'YOUR_BROWSERSTACK_ACCESS_KEY'],
        ['browser_version', 'latest'],
        ['os', 'Windows'],
        ['os_version', '11'],
        ['project_name', 'Aviasales Automation'],
        ['build_name', 'Flight Booking Test']
    ]
    
    # Add BrowserStack data
    for idx, row in enumerate(bs_data, start=2):
        ws_bs[f'A{idx}'] = row[0]
        ws_bs[f'B{idx}'] = row[1]
    
    # Adjust column widths
    ws_bs.column_dimensions['A'].width = 25
    ws_bs.column_dimensions['B'].width = 40
    
    # Add note cell
    note_row = len(bs_data) + 3
    ws_bs[f'A{note_row}'] = 'NOTE:'
    ws_bs[f'B{note_row}'] = 'Replace YOUR_BROWSERSTACK_USERNAME and YOUR_BROWSERSTACK_ACCESS_KEY with your actual BrowserStack credentials'
    ws_bs[f'A{note_row}'].font = Font(bold=True, color="FF0000")
    ws_bs[f'B{note_row}'].font = Font(italic=True)
    
    print(f"✓ Added {len(bs_data)} BrowserStack configuration fields")
    
    # =========================================
    # Save the workbook
    # =========================================
    filename = 'test_data.xlsx'
    
    try:
        wb.save(filename)
        print(f"\n{'='*60}")
        print(f"✓ SUCCESS: Excel file created: {filename}")
        print(f"{'='*60}")
        
        # Get file path
        file_path = os.path.abspath(filename)
        print(f"\nFile location: {file_path}")
        print(f"File size: {os.path.getsize(filename)} bytes")
        
        print("\n" + "="*60)
        print("NEXT STEPS:")
        print("="*60)
        print("1. Open test_data.xlsx and verify the data")
        print("2. If using BrowserStack:")
        print("   - Go to BrowserStack sheet")
        print("   - Replace YOUR_BROWSERSTACK_USERNAME with your username")
        print("   - Replace YOUR_BROWSERSTACK_ACCESS_KEY with your access key")
        print("3. Run the test:")
        print("   python test_flight_booking_with_excel.py")
        print("="*60)
        
        return True
        
    except Exception as e:
        print(f"\n❌ ERROR: Failed to create Excel file: {e}")
        return False


if __name__ == "__main__":
    # Check if openpyxl is installed
    try:
        import openpyxl
        print("✓ openpyxl library found\n")
    except ImportError:
        print("❌ ERROR: openpyxl library not found")
        print("\nPlease install it using:")
        print("pip install openpyxl")
        exit(1)
    
    # Create the Excel file
    success = create_test_data_excel()
    
    if success:
        print("\n✓ Excel file generation completed successfully!")
    else:
        print("\n❌ Excel file generation failed!")
        exit(1)
